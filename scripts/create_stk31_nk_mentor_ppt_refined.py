# -*- coding: utf-8 -*-
"""Build refined mentor PPT from frozen T/NK labels and refined CellChat results.
Does NOT overwrite the legacy PPT.
"""
from pathlib import Path
import csv
import subprocess

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.util import Inches, Pt

ROOT = Path(r"E:\ZHUWEIYU\Documents\R")
OUT_DIR = ROOT / "results" / "STK31_NK_mentor_ppt"
PNG_DIR = OUT_DIR / "png_refined"
PPTX_PATH = OUT_DIR / "STK31_NK_导师汇报_refined版_v2.pptx"

PDFTOPPM = Path(
    r"C:\Users\ZHUWEIYU\.cache\codex-runtimes\codex-primary-runtime\dependencies\native\poppler\Library\bin\pdftoppm.exe"
)

S1 = ROOT / "results" / "merged_tnk_refined_annotation"
S2 = ROOT / "results" / "merged_stk31_nk_refined_analysis"
S3 = ROOT / "results" / "merged_cellchat_refined_high_low_nk"
S4 = ROOT / "results" / "refined_sample_level_validation_v2"

FIGURES = {
    "umap_refined": S1 / "umap_refined_celltype.pdf",
    "dot_tnk": S1 / "dotplot_t_vs_nk_markers.pdf",
    "umap_analysis": S2 / "umap_analysis_celltype.pdf",
    "umap_rel": S2 / "umap_relationship_group.pdf",
    "nk_func": S2 / "nk_function_scores_by_sample.pdf",
    "cc_bubble": S3 / "cellchat_high_low_nk_bubble.pdf",
    "cc_heat": S3 / "cellchat_high_low_nk_heatmap.pdf",
    "cc_axis": S3 / "cellchat_mechanism_axis_high_low_barplot.pdf",
    "cc_circle": S3 / "cellchat_circle_count.pdf",
}


def ensure_dirs():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    PNG_DIR.mkdir(parents=True, exist_ok=True)


def pdf_to_png(name: str, pdf_path: Path) -> Path | None:
    if not pdf_path.exists():
        print(f"WARN missing figure: {pdf_path}")
        return None
    out_prefix = PNG_DIR / name
    png_path = PNG_DIR / f"{name}-1.png"
    if png_path.exists() and png_path.stat().st_mtime >= pdf_path.stat().st_mtime:
        return png_path
    if not PDFTOPPM.exists():
        print(f"WARN pdftoppm missing: {PDFTOPPM}")
        return None
    subprocess.run(
        [str(PDFTOPPM), "-png", "-singlefile", "-r", "200", str(pdf_path), str(out_prefix)],
        check=True,
    )
    single = PNG_DIR / f"{name}.png"
    if single.exists():
        single.replace(png_path)
    return png_path if png_path.exists() else None


def load_csv(path: Path):
    if not path.exists():
        return []
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def add_textbox(slide, text, left, top, width, height, font_size=16, bold=False, color=(30, 30, 30)):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = text
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(*color)
    return box


def add_title(slide, title, subtitle=None):
    add_textbox(slide, title, Inches(0.4), Inches(0.2), Inches(12.5), Inches(0.45), 22, True, (18, 64, 98))
    if subtitle:
        add_textbox(slide, subtitle, Inches(0.45), Inches(0.65), Inches(12.3), Inches(0.35), 11, False, (90, 90, 90))


def add_bullets(slide, bullets, left, top, width, height, font_size=14):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    for i, item in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item
        p.font.name = "Microsoft YaHei"
        p.font.size = Pt(font_size)
        p.font.color.rgb = RGBColor(40, 40, 40)
        p.space_after = Pt(5)
    return box


def add_note(slide, text):
    add_textbox(slide, text, Inches(0.45), Inches(6.85), Inches(12.3), Inches(0.3), 10, False, (120, 120, 120))


def add_image_fit(slide, image_path, left, top, width, height):
    if image_path is None or not Path(image_path).exists():
        add_textbox(slide, "[图缺失]", left, top, width, height, 14, False, (180, 80, 80))
        return None
    pic = slide.shapes.add_picture(str(image_path), left, top)
    scale = min(width / pic.width, height / pic.height)
    pic.width = int(pic.width * scale)
    pic.height = int(pic.height * scale)
    pic.left = int(left + (width - pic.width) / 2)
    pic.top = int(top + (height - pic.height) / 2)
    return pic


def add_table(slide, rows, headers, left, top, width, height, font_size=11):
    table = slide.shapes.add_table(len(rows) + 1, len(headers), left, top, width, height).table
    for j, header in enumerate(headers):
        cell = table.cell(0, j)
        cell.text = header
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor(18, 64, 98)
        for p in cell.text_frame.paragraphs:
            p.font.name = "Microsoft YaHei"
            p.font.size = Pt(font_size)
            p.font.bold = True
            p.font.color.rgb = RGBColor(255, 255, 255)
    for i, row in enumerate(rows, start=1):
        for j, value in enumerate(row):
            cell = table.cell(i, j)
            cell.text = str(value)
            for p in cell.text_frame.paragraphs:
                p.font.name = "Microsoft YaHei"
                p.font.size = Pt(font_size)
                p.font.color.rgb = RGBColor(35, 35, 35)
    return table


def build_ppt():
    ensure_dirs()
    imgs = {k: pdf_to_png(k, p) for k, p in FIGURES.items()}

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    # 1 title
    s = prs.slides.add_slide(blank)
    add_title(s, "STK31 与 NK 候选互作：refined 标签版汇报", "胆囊癌 scRNA-seq | 冻结 analysis_celltype | 2026-07-17")
    add_bullets(
        s,
        [
            "正式 NK 定义：analysis_celltype == \"NK_cell\"（冻结 refined 对象）",
            "高置信 NK：1,551（旧 cluster0/1 NK：7,419；保留约 20.8%）",
            "STK31-high 上皮：116（阳性率极低）",
            "结论边界：仅支持候选互作轴；不支持 STK31-high 特异增强；不支持因果",
        ],
        Inches(0.7),
        Inches(1.4),
        Inches(12),
        Inches(4.5),
        16,
    )
    add_note(s, "旧版 PPT 保留不覆盖；本文件为 refined 版。")

    # 2 unified conclusion
    s = prs.slides.add_slide(blank)
    add_title(s, "统一结论（必须遵守）")
    add_bullets(
        s,
        [
            "在严格拆分 T/NK 后，高置信 NK 从 7,419 降至 1,551。",
            "refined 分析仍提示 STK31-high 肿瘤上皮与 NK 之间存在候选互作轴，",
            "但当前 MHC-I、TIGIT/NECTIN、TGF-beta 和 NKG2D 的 high-vs-low 差异均为 similar，",
            "尚无证据支持 STK31-high 特异增强，也不支持因果关系。",
            "禁止：cluster 0/1=纯 NK；正式 NK=7419；特异增强 MHC-I；scTenifoldKnk 因果 KO。",
        ],
        Inches(0.6),
        Inches(1.3),
        Inches(12.2),
        Inches(5),
        16,
    )

    # 3 T/NK QC split table
    s = prs.slides.add_slide(blank)
    add_title(s, "T/NK 严格拆分 QC", "cluster 0/1 并非纯 NK；broad 原为 T_cell")
    add_table(
        s,
        [
            ["0", "877", "1602", "994", "375"],
            ["1", "669", "2274", "329", "299"],
        ],
        ["原 cluster", "NK", "T", "ambiguous", "contaminant"],
        Inches(0.6),
        Inches(1.2),
        Inches(7.5),
        Inches(1.6),
        13,
    )
    add_bullets(
        s,
        [
            "Stage1 pass：NK 的 CD3/TRAC 轴 < T；KLRD1/GNLY/NKG7/PRF1 > T",
            "高置信 NK 总数 1,551；tissue1/2/4/5 均有 refined NK",
            "ambiguous 不并入高置信 NK",
        ],
        Inches(0.6),
        Inches(3.2),
        Inches(12),
        Inches(2.8),
        15,
    )

    # 4 old vs new NK
    s = prs.slides.add_slide(blank)
    add_title(s, "新旧 NK 对比", "旧定义不可再用于主结论")
    add_table(
        s,
        [
            ["旧 NK（cluster 0/1）", "7419"],
            ["refined 高置信 NK", "1551"],
            ["交集", "1546"],
            ["旧有新无", "5873"],
            ["新有旧无", "5"],
            ["旧 NK 保留比例", "20.8%"],
        ],
        ["指标", "数值"],
        Inches(0.6),
        Inches(1.2),
        Inches(6.5),
        Inches(3.5),
        14,
    )
    add_bullets(
        s,
        [
            "旧有新无主要为 T 或 T/NK ambiguous",
            "基于旧 NK 的功能评分 / DE / CellChat 全部降级为 legacy",
            "主结果目录：merged_tnk_refined_annotation、merged_stk31_nk_refined_analysis、",
            "merged_cellchat_refined_high_low_nk、refined_sample_level_validation_v2",
            "v1 样本级目录已 superseded_method_issue，勿引用",
        ],
        Inches(7.4),
        Inches(1.3),
        Inches(5.4),
        Inches(4.5),
        13,
    )

    # 5 refined UMAP / dotplot
    s = prs.slides.add_slide(blank)
    add_title(s, "Refined 注释可视化", "仅高置信 NK 进入下游")
    add_image_fit(s, imgs.get("umap_refined"), Inches(0.4), Inches(1.1), Inches(6.2), Inches(5.4))
    add_image_fit(s, imgs.get("dot_tnk"), Inches(6.8), Inches(1.1), Inches(6.0), Inches(5.4))
    add_note(s, "左：refined_celltype UMAP；右：T vs NK marker DotPlot")

    # 6 relationship umap
    s = prs.slides.add_slide(blank)
    add_title(s, "STK31-high/low 上皮与 refined NK", "STK31-high 上皮 n=116")
    add_image_fit(s, imgs.get("umap_rel"), Inches(0.5), Inches(1.1), Inches(7.5), Inches(5.5))
    add_bullets(
        s,
        [
            "上皮 high/low 沿用原逻辑（cutoff=0 → counts>0）",
            "三种 STK31 定义敏感性均给出 high=116",
            "阳性率极低，解释需保守",
        ],
        Inches(8.2),
        Inches(1.5),
        Inches(4.6),
        Inches(4),
        14,
    )

    # 7 NK function
    s = prs.slides.add_slide(blank)
    add_title(s, "Refined NK 功能模块（按样本）", "仅 analysis_celltype==NK_cell")
    add_image_fit(s, imgs.get("nk_func"), Inches(0.5), Inches(1.2), Inches(12.2), Inches(5.2))
    add_note(s, "旧 followup 中基于 7419 NK 的功能评分已降级，不作为主结论。")

    # 8 CellChat high-low
    s = prs.slides.add_slide(blank)
    add_title(s, "CellChat high-vs-low → refined NK", "正式差分输出（非仅 high 组）")
    add_image_fit(s, imgs.get("cc_axis"), Inches(0.3), Inches(1.1), Inches(6.4), Inches(5.3))
    add_image_fit(s, imgs.get("cc_bubble"), Inches(6.8), Inches(1.1), Inches(6.1), Inches(5.3))
    add_note(s, "左：机制轴 high vs low；右：L-R bubble")

    # 9 MHC-I and axes
    s = prs.slides.add_slide(blank)
    add_title(s, "机制轴 high-vs-low 证据标签", "不得写成特异增强")
    add_table(
        s,
        [
            ["MHC-I/HLA", "0.1424", "0.1331", "+0.0093", "similar"],
            ["TIGIT/NECTIN", "—", "—", "小", "similar"],
            ["TGFb", "—", "—", "小", "similar"],
            ["NKG2D", "—", "—", "小", "similar"],
            ["IFNG/IFN", "0", "0", "0", "not_detected"],
        ],
        ["轴", "high→NK", "low→NK", "delta", "evidence_label"],
        Inches(0.5),
        Inches(1.2),
        Inches(12.2),
        Inches(3.2),
        13,
    )
    add_bullets(
        s,
        [
            "MHC-I：high 略高于 low，但 delta 很小，标签为 similar",
            "禁止表述：STK31-high 特异增强 MHC-I 抑制",
            "整体：存在候选互作轴，无特异增强与因果证据",
        ],
        Inches(0.6),
        Inches(4.7),
        Inches(12),
        Inches(1.8),
        14,
    )

    # 10 heatmap
    s = prs.slides.add_slide(blank)
    add_title(s, "High vs low → NK 互作热图（refined）")
    add_image_fit(s, imgs.get("cc_heat"), Inches(1.5), Inches(1.1), Inches(10), Inches(5.5))

    # 11 take-home
    s = prs.slides.add_slide(blank)
    add_title(s, "Take-home 与下一步")
    add_bullets(
        s,
        [
            "1) 标签已冻结：只用 refined analysis_celltype；NK=1551",
            "2) 旧 NK=7419 相关结论全部 legacy",
            "3) refined CellChat：主轴 similar / IFNG not_detected（主结论依据）",
            "4) 样本级 v2：六轴均为 exploratory_mixed；direction_consistent=FALSE",
            "   LOSO majority stable ≠ 四样本方向一致；within-sample bootstrap 不升级证据",
            "5) STK31 三定义敏感性 non-informative（均退化为 detected vs not）",
            "6) 后续：批次整合、恶性上皮确认、外部验证、实验",
            "7) 表述红线：候选互作轴；非特异增强；非因果",
        ],
        Inches(0.7),
        Inches(1.3),
        Inches(12),
        Inches(5),
        16,
    )

    prs.save(PPTX_PATH)
    print(f"Wrote {PPTX_PATH}")
    return PPTX_PATH


if __name__ == "__main__":
    build_ppt()
